"""Branded file generators: Excel (openpyxl) and PDF (reportlab) for dotindot exports."""
import os
import re
from io import BytesIO
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ORANGE = "F26B21"
ORANGE_HEX = "#F26B21"

# ---- brand (from the official logo) ----
BRAND_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "brand")
BRAND_WORDMARK = os.path.join(BRAND_DIR, "wordmark-color.png")     # 1600 × 307, transparent
BRAND_WORDMARK_WHITE = os.path.join(BRAND_DIR, "wordmark-white.png")
BRAND_MARK = os.path.join(BRAND_DIR, "mark-color.png")             # the "D", transparent
BRAND_WORDMARK_XLSX = os.path.join(BRAND_DIR, "wordmark-xlsx.png")  # 270 px wide, for spreadsheets
WORDMARK_RATIO = 396 / 76          # width / height
MARK_RATIO = 140.5 / 179
BRAND_FROM, BRAND_TO = "#FE7A18", "#FFAD42"   # wordmark gradient

# Register a unicode font so the ₹ glyph renders in PDFs; fall back to Helvetica + "Rs"
try:
    pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
    FONT, BOLD_FONT, RUPEE = "DejaVu", "DejaVu-Bold", "\u20b9"
except Exception:
    FONT, BOLD_FONT, RUPEE = "Helvetica", "Helvetica-Bold", "Rs "


def inr(n):
    """Indian-grouped currency string: 185000 -> ₹1,85,000"""
    try:
        n = float(n)
    except (TypeError, ValueError):
        return str(n or "")
    neg = n < 0
    digits = str(int(round(abs(n))))
    if len(digits) > 3:
        head, tail = digits[:-3], digits[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        digits = ",".join(groups + [tail])
    return ("-" if neg else "") + RUPEE + digits


def _stamp():
    return datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return str(v)


# ---------------- Excel ----------------
def _xlsx_logo(ws):
    """dotindot. wordmark in the top-left corner (row 1)."""
    ws.row_dimensions[1].height = 30
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(BRAND_WORDMARK_XLSX)
        img.height, img.width = 24, round(24 * WORDMARK_RATIO)
        ws.add_image(img, "A1")
    except Exception:  # no Pillow / missing asset → plain text fallback
        ws.cell(1, 1, "dotindot.").font = Font(bold=True, size=14, color="FE7A18")


def build_xlsx(title, sections, summary=None, generated_by=""):
    """sections: [{"title": str, "columns": [(key,label,fmt)], "rows": [dict]}]; fmt in text|money|num"""
    wb = Workbook()
    wb.remove(wb.active)

    if summary:
        ws = wb.create_sheet("Summary")
        _xlsx_logo(ws)
        ws.merge_cells("A2:C2")
        c = ws.cell(2, 1, title)
        c.font = Font(bold=True, size=14, color="111827")
        ws.cell(3, 1, f"Generated {_stamp()} by {generated_by}").font = Font(size=9, color="888888")
        r = 5
        for item in summary:
            ws.cell(r, 1, item["label"]).font = Font(bold=True, size=10)
            vc = ws.cell(r, 2)
            if item.get("money"):
                vc.value = float(item["value"] or 0)
                vc.number_format = '"₹"#,##0'
            else:
                vc.value = item["value"]
            vc.font = Font(size=10)
            r += 1
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 22

    for idx, sec in enumerate(sections):
        raw = sec.get("title") or title or f"Sheet{idx + 1}"
        name = re.sub(r"[\\/*?:\[\]]", "", raw)[:31] or f"Sheet{idx + 1}"
        if name in wb.sheetnames:
            name = f"{name[:28]}_{idx + 1}"
        ws = wb.create_sheet(name)
        cols = sec["columns"]
        ncol = max(len(cols), 1)
        _xlsx_logo(ws)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        c = ws.cell(2, 1, title + (f" — {sec.get('title')}" if sec.get("title") else ""))
        c.font = Font(bold=True, size=13, color="111827")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncol)
        ws.cell(3, 1, f"Generated {_stamp()} by {generated_by} · {len(sec['rows'])} rows").font = Font(size=9, color="888888")

        hr = 5
        widths = []
        for j, (key, label, fmt) in enumerate(cols, start=1):
            cell = ws.cell(hr, j, label)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor=ORANGE)
            cell.alignment = Alignment(horizontal="right" if fmt in ("money", "num") else "left")
            widths.append(len(label) + 4)
        for i, row in enumerate(sec["rows"], start=hr + 1):
            for j, (key, label, fmt) in enumerate(cols, start=1):
                v = row.get(key)
                cell = ws.cell(i, j)
                if fmt == "money":
                    cell.value = float(v or 0)
                    cell.number_format = '"₹"#,##0'
                elif fmt == "num":
                    cell.value = v if isinstance(v, (int, float)) else _cell_text(v)
                else:
                    cell.value = _cell_text(v)
                widths[j - 1] = min(50, max(widths[j - 1], len(_cell_text(v)) + 2))
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(hr + 1, 1)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------- PDF ----------------
def _gradient_rule(canvas, x, y, w, h):
    """Thin brand-gradient bar (left → right)."""
    canvas.saveState()
    try:
        path = canvas.beginPath()
        path.rect(x, y, w, h)
        canvas.clipPath(path, stroke=0, fill=0)
        canvas.linearGradient(x, y, x + w, y, (colors.HexColor(BRAND_FROM), colors.HexColor(BRAND_TO)), extend=False)
    except Exception:
        canvas.setFillColor(colors.HexColor(BRAND_FROM))
        canvas.rect(x, y, w, h, stroke=0, fill=1)
    canvas.restoreState()


def draw_wordmark(canvas, x, y, height, white=False):
    """dotindot. wordmark with its bottom-left corner at (x, y). Returns the drawn width."""
    width = height * WORDMARK_RATIO
    path = BRAND_WORDMARK_WHITE if white else BRAND_WORDMARK
    try:
        canvas.drawImage(path, x, y, width=width, height=height, mask="auto")
    except Exception:
        canvas.setFillColor(colors.white if white else colors.HexColor(BRAND_FROM))
        canvas.setFont(BOLD_FONT, height * 1.15)
        canvas.drawString(x, y + height * 0.12, "dotindot.")
    return width


def _page_decorator(title, generated_by):
    stamp = _stamp()

    def draw(canvas, doc):
        w, h = doc.pagesize
        canvas.saveState()
        # header: full-colour logo on white, title on the right, brand gradient rule underneath
        draw_wordmark(canvas, 12 * mm, h - 12.6 * mm, 6.4 * mm)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.setFont(FONT, 8.5)
        canvas.drawRightString(w - 12 * mm, h - 10.6 * mm, title[:90])
        _gradient_rule(canvas, 0, h - 16 * mm, w, 1.1 * mm)
        # footer: small mark + who/when, page number
        try:
            canvas.drawImage(BRAND_MARK, 12 * mm, 6.9 * mm, width=3.4 * mm * MARK_RATIO, height=3.4 * mm, mask="auto")
            text_x = 12 * mm + 3.4 * mm * MARK_RATIO + 2 * mm
        except Exception:
            text_x = 12 * mm
        canvas.setFillColor(colors.HexColor("#9CA3AF"))
        canvas.setFont(FONT, 7.5)
        canvas.drawString(text_x, 8 * mm, f"Generated by {generated_by} · {stamp} · dotindot Ops Platform")
        canvas.drawRightString(w - 12 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    return draw


def build_pdf(title, subtitle, sections, summary=None, generated_by="", max_rows=600):
    ncols_max = max((len(s["columns"]) for s in sections), default=1)
    pagesize = landscape(A4) if ncols_max > 6 else A4
    bio = BytesIO()
    doc = SimpleDocTemplate(
        bio, pagesize=pagesize, leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=22 * mm, bottomMargin=16 * mm, title=f"dotindot — {title}",
    )
    usable = pagesize[0] - 24 * mm

    h1 = ParagraphStyle("h1", fontName=BOLD_FONT, fontSize=15, leading=19, textColor=colors.HexColor("#111827"), spaceAfter=2)
    sub = ParagraphStyle("sub", fontName=FONT, fontSize=9, leading=12, textColor=colors.HexColor("#6B7280"), spaceAfter=8)
    h2 = ParagraphStyle("h2", fontName=BOLD_FONT, fontSize=11, textColor=colors.HexColor(ORANGE_HEX), spaceBefore=10, spaceAfter=4)
    cellstyle = ParagraphStyle("cell", fontName=FONT, fontSize=7.5, textColor=colors.HexColor("#1F2937"), leading=9)

    story = [Paragraph(title, h1)]
    if subtitle:
        story.append(Paragraph(subtitle, sub))

    if summary:
        chips = []
        for item in summary:
            val = inr(item["value"]) if item.get("money") else _cell_text(item["value"])
            chips.append([Paragraph(f"<b>{val}</b>", ParagraphStyle("v", fontName=BOLD_FONT, fontSize=11, textColor=colors.HexColor(ORANGE_HEX))),
                          ])
        data = [[Paragraph(_cell_text(i["label"]), ParagraphStyle("l", fontName=FONT, fontSize=7, textColor=colors.HexColor("#6B7280"))) for i in summary],
                [Paragraph(f"<b>{inr(i['value']) if i.get('money') else _cell_text(i['value'])}</b>",
                           ParagraphStyle("v", fontName=BOLD_FONT, fontSize=10.5, textColor=colors.HexColor("#111827"))) for i in summary]]
        t = Table(data, colWidths=[usable / len(summary)] * len(summary))
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF7ED")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#FDBA74")),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#FED7AA")),
            ("TOPPADDING", (0, 0), (-1, 0), 5), ("BOTTOMPADDING", (0, 1), (-1, 1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([t, Spacer(1, 6)])

    for sec in sections:
        cols = sec["columns"]
        rows = sec["rows"][:max_rows]
        if sec.get("title"):
            story.append(Paragraph(sec["title"], h2))
        header = [Paragraph(f"<b>{label}</b>", ParagraphStyle("hd", fontName=BOLD_FONT, fontSize=7.5, textColor=colors.white)) for _, label, _ in cols]
        data = [header]
        for row in rows:
            line = []
            for key, label, fmt in cols:
                v = row.get(key)
                txt = inr(v) if fmt == "money" else _cell_text(v)
                if len(txt) > 70:
                    txt = txt[:67] + "..."
                line.append(Paragraph(txt.replace("&", "&amp;").replace("<", "&lt;"), cellstyle))
            data.append(line)
        money_idx = [i for i, (_, _, f) in enumerate(cols) if f in ("money", "num")]
        weights = [1.6 if f == "text" and k in ("description", "name", "title", "notes", "entity_name") else 1 for k, _, f in cols]
        total_w = sum(weights)
        col_widths = [usable * w / total_w for w in weights]
        t = Table(data, colWidths=col_widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(ORANGE_HEX)),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF9F3")]),
        ]
        for i in money_idx:
            style.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
        t.setStyle(TableStyle(style))
        story.append(t)
        if len(sec["rows"]) > max_rows:
            story.append(Paragraph(f"Showing first {max_rows} of {len(sec['rows'])} rows — download the Excel export for the complete data set.", sub))
        story.append(Spacer(1, 4))

    deco = _page_decorator(title, generated_by)
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return bio.getvalue()
